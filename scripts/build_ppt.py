"""備審 PPT 產生器 — Agent 請修改此檔案

使用方法：
  1. Agent 讀取 ../profile.md 取得使用者資料
  2. 將下方所有 {{PLACEHOLDER}} 取代為實際內容
  3. 每段新增文字上方加上 # TRUTH: profile.md → 章節代號
  4. 在 main() 中啟用需要產出的類別
  5. 執行 python build_ppt.py
"""

import os, sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from pptx_engine import *
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

# ============================================================
# DIRECTORIES — Agent 請勿修改
# ============================================================
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
IMAGES_DIR = os.path.join(ROOT, "images")
OUTPUT_DIR = os.path.join(ROOT, "output")
os.makedirs(OUTPUT_DIR, exist_ok=True)

def img(filename):
    return os.path.join(IMAGES_DIR, filename) if filename else ""

# ============================================================
# 個人資料（Agent 請從 profile.md 填入，不得額外發明）
# ============================================================
# TRUTH: profile.md → 1.1
STUDENT_NAME_ZH = "{{NAME_ZH}}"
# TRUTH: profile.md → 1.2
STUDENT_NAME_EN = "{{NAME_EN}}"
# TRUTH: profile.md → 1.3
STUDENT_SCHOOL = "{{SCHOOL}}"
# TRUTH: profile.md → 1.4
STUDENT_DEPT = "{{DEPARTMENT}}"
# TRUTH: profile.md → 1.6
STUDENT_EMAIL = "{{EMAIL}}"
# TRUTH: profile.md → 1.7
STUDENT_GITHUB = "{{GITHUB}}"
# TRUTH: profile.md → 1.9
HEADSHOT_FILE = img("{{HEADSHOT_FILENAME}}")
FOOTER_TEXT = f"{STUDENT_NAME_ZH} | {STUDENT_SCHOOL}{STUDENT_DEPT}"

# ============================================================
# 目標校系
# ============================================================
# TRUTH: profile.md → 2.1
TARGET_SCHOOL = "{{TARGET_SCHOOL}}"
# TRUTH: profile.md → 2.2
TARGET_DEPT = "{{TARGET_DEPARTMENT}}"
# TRUTH: profile.md → 2.3
TARGET_GROUP = "{{TARGET_GROUP}}"

# ============================================================
# 主打專題資料（B-1a）
# ============================================================
# TRUTH: profile.md → 3.1
PROJECT_NAME = "{{PROJECT_NAME_ZH}}"
# TRUTH: profile.md → 3.3
PROJECT_DATES = "{{PROJECT_START}} ~ {{PROJECT_END}}"
# TRUTH: profile.md → 3.4
PROJECT_DURATION = "{{PROJECT_DURATION}}"
# TRUTH: profile.md → 3.5
PROJECT_TEAM_SIZE = "{{PROJECT_TEAM_SIZE}}"
# TRUTH: profile.md → 3.6
PROJECT_ROLE = "{{PROJECT_ROLE}}"
# TRUTH: profile.md → 3.7
TEAM_MEMBERS = "{{TEAM_MEMBER_1}}：{{TEAM_MEMBER_1_ROLE}}；{{TEAM_MEMBER_2}}：{{TEAM_MEMBER_2_ROLE}}"
# TRUTH: profile.md → 4.1
HARDWARE = ["{{HW_PART_1}}", "{{HW_PART_2}}", "{{HW_PART_3}}", "{{HW_PART_4}}", "{{HW_PART_5}}"]
# TRUTH: profile.md → 4.2
SOFTWARE = ["{{SW_TOOL_1}}", "{{SW_TOOL_2}}", "{{SW_TOOL_3}}"]
# TRUTH: profile.md → 4.3
PROTOCOL = "{{COMM_PROTOCOL}}"
# TRUTH: profile.md → 4.5
GIT_COMMITS = "{{GIT_COMMITS}}"
# TRUTH: profile.md → 4.6
GIT_BRANCH = "{{GIT_BRANCH_STRATEGY}}"
# TRUTH: profile.md → 5.1
MOTIVATION = "{{MOTIVATION_ORIGINAL}}"
# TRUTH: profile.md → 5.2
PROBLEM = "{{PROBLEM_STATEMENT}}"
# TRUTH: profile.md → 5.4
SURVEY_N = "{{SURVEY_SAMPLE_SIZE}}"
# TRUTH: profile.md → 5.5
SURVEY_RESULTS = ["{{SURVEY_RESULT_1}}", "{{SURVEY_RESULT_2}}", "{{SURVEY_RESULT_3}}", "{{SURVEY_RESULT_4}}"]
# TRUTH: profile.md → 6.1
CONTRIB_SUMMARY = "{{CONTRIBUTION_SUMMARY}}"
# TRUTH: profile.md → 6.2
HW_CONTRIB = ["{{HW_CONTRIB_1}}", "{{HW_CONTRIB_2}}", "{{HW_CONTRIB_3}}"]
# TRUTH: profile.md → 6.3
SW_CONTRIB = ["{{SW_CONTRIB_1}}", "{{SW_CONTRIB_2}}", "{{SW_CONTRIB_3}}"]
# TRUTH: profile.md → 6.4
MECH_CONTRIB = ["{{MECH_CONTRIB_1}}", "{{MECH_CONTRIB_2}}"]
# TRUTH: profile.md → 7.2
DIFFICULTY_STORY = "{{DIFFICULTY_1_STORY}}"
# TRUTH: profile.md → 7.4
MEMORABLE_NIGHT = "{{MEMORABLE_NIGHT_STORY}}"
# TRUTH: profile.md → 7.6
FAILURE_STORY = "{{FAILURE_STORY}}"
# TRUTH: profile.md → 7.7
TOLERANCE_STORY = "{{TOLERANCE_STORY}}"
# TRUTH: profile.md → 7.8
GIT_STORY = "{{GIT_DISASTER_STORY}}"
# TRUTH: profile.md → 8.1
LESSONS = ["{{LESSON_1}}", "{{LESSON_2}}", "{{LESSON_3}}"]
# TRUTH: profile.md → 8.2
DEMONSTRATED_SKILLS = "{{DEMONSTRATED_SKILLS}}"
# TRUTH: profile.md → 9.1
PROJECT_PHOTOS = ["{{PHOTO_1_FILENAME}}", "{{PHOTO_2_FILENAME}}", "{{PHOTO_3_FILENAME}}", "{{PHOTO_4_FILENAME}}", "{{PHOTO_5_FILENAME}}", "{{PHOTO_6_FILENAME}}"]
# TRUTH: profile.md → 9.2
PROJECT_PDF = "{{PROJECT_PDF}}"

# ============================================================
# 競賽獎項（C-5）
# ============================================================
# TRUTH: profile.md → 10.1
AWARD_1_NAME = "{{AWARD_1_NAME}}"
# TRUTH: profile.md → 10.2
AWARD_1_ORG = "{{AWARD_1_ORGANIZER}}"
# TRUTH: profile.md → 10.4
AWARD_1_RANK = "{{AWARD_1_RANK}}"
# TRUTH: profile.md → 10.5
AWARD_1_SCALE = "{{AWARD_1_SCALE}}"
# TRUTH: profile.md → 10.7
AWARD_1_CERT = img("{{AWARD_1_CERT_FILENAME}}")
# （如有第二項競賽，Agent 自行新增 AWARD_2 系列變數）

# ============================================================
# 證照（C-7）
# ============================================================
# TRUTH: profile.md → 11.1
CERT_1_NAME = "{{CERT_1_NAME}}"
# TRUTH: profile.md → 11.2
CERT_1_ORG = "{{CERT_1_ISSUER}}"
# TRUTH: profile.md → 11.5
CERT_1_FILE = img("{{CERT_1_FILENAME}}")
# TRUTH: profile.md → 11.1 (第二張)
CERT_2_NAME = "{{CERT_2_NAME}}"
# TRUTH: profile.md → 11.2
CERT_2_ORG = "{{CERT_2_ISSUER}}"
# TRUTH: profile.md → 11.5
CERT_2_FILE = img("{{CERT_2_FILENAME}}")

# ============================================================
# 課外活動（C-8）
# ============================================================
# TRUTH: profile.md → 12.1
ACTIVITIES = [
    {"name": "{{ACTIVITY_1_NAME}}", "org": "{{ACTIVITY_1_ORGANIZER}}", "hours": "{{ACTIVITY_1_DATE_HOURS}}", "learn": "{{ACTIVITY_1_LEARNING}}", "cert": img("{{ACTIVITY_1_CERT_FILENAME}}")},
]

# ============================================================
# 就讀動機（D-2）
# ============================================================
# TRUTH: profile.md → 13.1
MOTIVATION_STATEMENT = "{{MOTIVATION_STATEMENT}}"
# TRUTH: profile.md → 13.2
KEY_LABELS = ["{{KEY_LABEL_1}}", "{{KEY_LABEL_2}}", "{{KEY_LABEL_3}}"]
# TRUTH: profile.md → 13.3
COLLEGE_PLAN = "{{COLLEGE_PLAN}}"
# TRUTH: profile.md → 13.4
LONG_TERM_GOAL = "{{LONG_TERM_GOAL}}"
# TRUTH: profile.md → 13.5
MESSAGE_TO_PROF = "{{MESSAGE_TO_PROFESSOR}}"

# ============================================================
# 其他作品（D-3）
# ============================================================
# TRUTH: profile.md → 14.1
OTHER_PROJECTS = [
    {"name": "{{OTHER_PROJECT_NAME}}", "tech": "{{OTHER_PROJECT_TECH}}", "link": "{{OTHER_PROJECT_LINK}}", "contrib": "{{OTHER_PROJECT_CONTRIBUTION}}"},
]

# ============================================================
# 實習科目（B-1b）— Agent 如需要請從此填入
# ============================================================
# TRUTH: profile.md → 15.1
INTERNSHIP_COURSES = ["{{INTERNSHIP_COURSE_1}}", "{{INTERNSHIP_COURSE_2}}"]
# TRUTH: profile.md → 15.2
INTERNSHIP_WORKS = ["{{INTERNSHIP_WORK_1}}", "{{INTERNSHIP_WORK_2}}"]
# TRUTH: profile.md → 15.3
INTERNSHIP_LESSON = "{{INTERNSHIP_LESSON}}"
# TRUTH: profile.md → 15.5
INTERNSHIP_PHOTOS = [img("{{INTERNSHIP_PHOTO_1}}"), img("{{INTERNSHIP_PHOTO_2}}")]

# ============================================================
# 手繪製圖（B-2a）— Agent 如需要請從此填入
# ============================================================
# TRUTH: profile.md → 16.1
HANDDRAFT_WORKS = ["{{HANDDRAFT_WORK_1}}", "{{HANDDRAFT_WORK_2}}", "{{HANDDRAFT_WORK_3}}"]
# TRUTH: profile.md → 16.2
HANDDRAFT_REFLECTION = "{{HANDDRAFT_REFLECTION}}"
# TRUTH: profile.md → 16.3
HANDDRAFT_PHOTOS = [img("{{HANDDRAFT_PHOTO_1}}"), img("{{HANDDRAFT_PHOTO_2}}")]

# ============================================================
# 電腦輔助製圖（B-2b）
# ============================================================
# TRUTH: profile.md → 17.1
CAD_WORKS = ["{{CAD_WORK_1}}", "{{CAD_WORK_2}}"]
# TRUTH: profile.md → 17.3
CAD_REFLECTION = "{{CAD_REFLECTION}}"

# ============================================================
# 3D 建模（B-2c）
# ============================================================
# TRUTH: profile.md → 18.1
MODEL_WORKS = ["{{MODEL_WORK_1}}", "{{MODEL_WORK_2}}"]

# ============================================================
# 彈性學習（C-1）
# ============================================================
# TRUTH: profile.md → 19.1
FLEXIBLE_LEARNING = "{{FLEXIBLE_LEARNING_NAME}}"
# TRUTH: profile.md → 19.3
FLEXIBLE_REFLECTION = "{{FLEXIBLE_REFLECTION}}"

# ============================================================
# 多元表現綜整心得（D-1）
# ============================================================
# TRUTH: profile.md → 20.1
SUMMARY_TEXT = "{{SUMMARY_TEXT}}"


# ====================================================================
#  PPT 產生函式 — Agent 請逐一修改內容，加入 TRUTH 標記
# ====================================================================

def build_project_ppt():
    """B-1a: 專題實作"""
    prs = new_presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_cover(s, f"B-1 專題實作\n{PROJECT_NAME}")

    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_toc(s, [
        ("01", "研究動機與目的", ""),
        ("02", "硬體架構與通訊", ""),
        ("03", "機械機構設計", ""),
        ("04", "App 功能頁面", ""),
        ("05", "團隊分工與貢獻", ""),
        ("06", "開發規模", ""),
        ("07", "反思與學習", ""),
    ], highlight_text=PROJECT_NAME, photo_path=img(PROJECT_PHOTOS[0]) if PROJECT_PHOTOS and PROJECT_PHOTOS[0] else None)

    # 研究動機
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(s, "研究動機與目的")
    lines = [MOTIVATION, "", PROBLEM]
    if SURVEY_N:
        lines += ["", f"問卷調查（n={SURVEY_N}）："]
        for r in SURVEY_RESULTS:
            if r and r.startswith("{{") == False:
                lines.append(f"• {r}")
    add_multiline_text(s, 0.8, 1.6, 11.5, 5.5, lines, font_size=16)

    # 硬體架構
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(s, "硬體架構與通訊")
    hw_lines = []
    for h in HARDWARE:
        if h and h.startswith("{{") == False:
            hw_lines.append(f"• {h}")
    if PROTOCOL and PROTOCOL.startswith("{{") == False:
        hw_lines += ["", f"通訊協定：{PROTOCOL}"]
    add_multiline_text(s, 0.8, 1.6, 6.0, 4.0, hw_lines, font_size=18)
    for i, p in enumerate(PROJECT_PHOTOS):
        if p and p.startswith("{{") == False:
            x = 6.6 + (i % 2) * 3.3
            y = 1.6 + (i // 2) * 2.8
            add_image(s, img(p), x, y, 3.0)

    # 團隊分工
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(s, "團隊分工與貢獻")
    team_lines = [f"團隊人數：{PROJECT_TEAM_SIZE}", f"角色：{PROJECT_ROLE}", ""]
    for c in HW_CONTRIB:
        if c and c.startswith("{{") == False:
            team_lines.append(f"• {c}")
    for c in SW_CONTRIB:
        if c and c.startswith("{{") == False:
            team_lines.append(f"• {c}")
    add_multiline_text(s, 0.8, 1.6, 11.5, 5, team_lines, font_size=16)

    # 開發規模
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(s, "開發規模")
    git_lines = [f"開發期間：{PROJECT_DATES}", f"時長：{PROJECT_DURATION}"]
    if GIT_COMMITS and GIT_COMMITS.startswith("{{") == False:
        git_lines.append(f"總 commits：{GIT_COMMITS}")
    if GIT_BRANCH and GIT_BRANCH.startswith("{{") == False:
        git_lines.append(f"分支策略：{GIT_BRANCH}")
    add_multiline_text(s, 0.8, 1.6, 11.5, 4, git_lines, font_size=18)

    # 困難故事
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(s, "開發過程的挑戰")
    story_parts = []
    if DIFFICULTY_STORY and DIFFICULTY_STORY.startswith("{{") == False:
        story_parts.append(DIFFICULTY_STORY)
    if MEMORABLE_NIGHT and MEMORABLE_NIGHT.startswith("{{") == False:
        story_parts.append("", MEMORABLE_NIGHT)
    add_multiline_text(s, 0.8, 1.6, 11.5, 5.5, story_parts, font_size=16)

    # 反思
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(s, "反思與學習")
    lesson_lines = []
    for l in LESSONS:
        if l and l.startswith("{{") == False:
            lesson_lines.append(f"• {l}")
    add_multiline_text(s, 0.8, 1.6, 11.5, 5, lesson_lines, font_size=16)

    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_ending(s)
    add_footer(prs, FOOTER_TEXT)
    path = os.path.join(OUTPUT_DIR, "B-1a_專題實作.pptx")
    prs.save(path)
    print(f"[OK] B-1a: {os.path.getsize(path):,} bytes")


def build_internship_ppt():
    """B-1b: 實習科目學習成果（車床/銑床/鉗工/焊接/電學等）"""
    prs = new_presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_cover(s, "B-1 實習科目學習成果")

    s = prs.slides.add_slide(prs.slide_layouts[6])
    toc_items = []
    for i, c in enumerate(INTERNSHIP_COURSES):
        if c and c.startswith("{{") == False:
            toc_items.append((f"{i+1:02d}", c, ""))
    add_toc(s, toc_items, highlight_text="實習歷程與反思")

    for i, c in enumerate(INTERNSHIP_COURSES):
        if c and c.startswith("{{") == False:
            s = prs.slides.add_slide(prs.slide_layouts[6])
            add_title(s, c)
            work = INTERNSHIP_WORKS[i] if i < len(INTERNSHIP_WORKS) else ""
            lines = []
            if work and work.startswith("{{") == False:
                lines.append(f"作品：{work}")
            lines.append("")
            if INTERNSHIP_LESSON and INTERNSHIP_LESSON.startswith("{{") == False:
                lines.append(INTERNSHIP_LESSON)
            add_multiline_text(s, 0.8, 1.6, 6, 4, lines, font_size=18)
            if i < len(INTERNSHIP_PHOTOS) and INTERNSHIP_PHOTOS[i]:
                add_image(s, INTERNSHIP_PHOTOS[i], 7.5, 1.6, 4.5)

    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_ending(s)
    add_footer(prs, FOOTER_TEXT)
    path = os.path.join(OUTPUT_DIR, "B-1b_實習科目.pptx")
    prs.save(path)
    print(f"[OK] B-1b: {os.path.getsize(path):,} bytes")


def build_handdraft_ppt():
    """B-2a: 手繪製圖"""
    prs = new_presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_cover(s, "B-2 其他課程學習成果", "手繪製圖")

    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_toc(s, [
        ("01", "手繪製圖學習歷程", ""),
        ("02", "作品展示", ""),
        ("03", "心得反思", ""),
    ])

    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(s, "手繪製圖學習歷程")
    add_multiline_text(s, 0.8, 1.6, 11.5, 5.5, [
        "訓練空間想像力與圖面表達能力",
        "• 三視圖、剖面圖、等角圖繪製",
        "• 尺寸標註與公差配合",
        "• 線條練習與圖面配置",
    ], font_size=18)

    if HANDDRAFT_PHOTOS:
        for i, p in enumerate(HANDDRAFT_PHOTOS):
            if p:
                s = prs.slides.add_slide(prs.slide_layouts[6])
                add_title(s, "手繪作品展示")
                add_image(s, p, 0.8, 1.6, 11.5)

    if HANDDRAFT_REFLECTION and HANDDRAFT_REFLECTION.startswith("{{") == False:
        s = prs.slides.add_slide(prs.slide_layouts[6])
        add_title(s, "心得反思")
        add_multiline_text(s, 0.8, 1.6, 11.5, 5.5, [HANDDRAFT_REFLECTION], font_size=18)

    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_ending(s)
    add_footer(prs, FOOTER_TEXT)
    path = os.path.join(OUTPUT_DIR, "B-2a_手繪製圖.pptx")
    prs.save(path)
    print(f"[OK] B-2a: {os.path.getsize(path):,} bytes")


def build_cad_ppt():
    """B-2b: 電腦輔助製圖"""
    prs = new_presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_cover(s, "B-2 其他課程學習成果", "電腦輔助製圖")

    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_toc(s, [
        ("01", "電腦輔助製圖學習歷程", ""),
        ("02", "作品展示", ""),
        ("03", "心得反思", ""),
    ])

    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(s, "電腦輔助製圖學習歷程")
    add_multiline_text(s, 0.8, 1.6, 11.5, 5.5, [
        "從 AutoCAD 2D 到 Inventor 3D 建模",
        "• 圖層管理、線型設定、圖塊應用",
        "• 3D 實體建模、組合圖、爆炸圖",
        "• 丙級檢定練習",
    ], font_size=18)

    if CAD_REFLECTION and CAD_REFLECTION.startswith("{{") == False:
        s = prs.slides.add_slide(prs.slide_layouts[6])
        add_title(s, "心得反思")
        add_multiline_text(s, 0.8, 1.6, 11.5, 5.5, [CAD_REFLECTION], font_size=18)

    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_ending(s)
    add_footer(prs, FOOTER_TEXT)
    path = os.path.join(OUTPUT_DIR, "B-2b_電腦輔助製圖.pptx")
    prs.save(path)
    print(f"[OK] B-2b: {os.path.getsize(path):,} bytes")


def build_3dmodel_ppt():
    """B-2c: 3D 建模作品"""
    prs = new_presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_cover(s, "B-2 其他課程學習成果", "3D 建模作品")

    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_toc(s, [
        ("01", "3D 建模作品", ""),
    ])

    for i, m in enumerate(MODEL_WORKS):
        if m and m.startswith("{{") == False:
            s = prs.slides.add_slide(prs.slide_layouts[6])
            add_title(s, m)
            if PROJECT_PHOTOS and i < len(PROJECT_PHOTOS):
                add_image(s, img(PROJECT_PHOTOS[i]), 0.8, 1.6, 5.0)

    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_ending(s)
    add_footer(prs, FOOTER_TEXT)
    path = os.path.join(OUTPUT_DIR, "B-2c_3D建模.pptx")
    prs.save(path)
    print(f"[OK] B-2c: {os.path.getsize(path):,} bytes")


def build_flexible_ppt():
    """C-1: 彈性學習時間成果"""
    prs = new_presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_cover(s, "C-1 彈性學習時間成果")

    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(s, FLEXIBLE_LEARNING if FLEXIBLE_LEARNING.startswith("{{") == False else "彈性學習")
    lines = []
    if FLEXIBLE_REFLECTION and FLEXIBLE_REFLECTION.startswith("{{") == False:
        lines.append(FLEXIBLE_REFLECTION)
    add_multiline_text(s, 0.8, 1.6, 11.5, 5.5, lines, font_size=18)

    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_ending(s)
    add_footer(prs, FOOTER_TEXT)
    path = os.path.join(OUTPUT_DIR, "C-1_彈性學習.pptx")
    prs.save(path)
    print(f"[OK] C-1: {os.path.getsize(path):,} bytes")


def build_competition_ppt():
    """C-5: 競賽表現"""
    prs = new_presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_cover(s, "C-5 競賽表現")

    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_toc(s, [("01", AWARD_1_NAME, AWARD_1_RANK)])

    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(s, AWARD_1_NAME)
    lines = [f"主辦單位：{AWARD_1_ORG}", f"獎項：{AWARD_1_RANK}"]
    if AWARD_1_SCALE and AWARD_1_SCALE.startswith("{{") == False:
        lines.append(f"參賽規模：{AWARD_1_SCALE}")
    add_multiline_text(s, 0.8, 1.6, 6, 4, lines, font_size=18)
    if AWARD_1_CERT:
        add_image(s, AWARD_1_CERT, 7.5, 1.6, 4)

    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_ending(s)
    add_footer(prs, FOOTER_TEXT)
    path = os.path.join(OUTPUT_DIR, "C-5_競賽表現.pptx")
    prs.save(path)
    print(f"[OK] C-5: {os.path.getsize(path):,} bytes")


def build_cert_pro_ppt():
    """C-7a: 專業檢定證照"""
    prs = new_presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_cover(s, "C-7 檢定證照", "專業技術士證照")

    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_toc(s, [("01", CERT_1_NAME, CERT_1_ORG), ("02", CERT_2_NAME, CERT_2_ORG)])

    for name, org, cert_file in [(CERT_1_NAME, CERT_1_ORG, CERT_1_FILE), (CERT_2_NAME, CERT_2_ORG, CERT_2_FILE)]:
        if name and name.startswith("{{") == False:
            s = prs.slides.add_slide(prs.slide_layouts[6])
            add_title(s, name)
            add_multiline_text(s, 0.8, 1.6, 6, 4, [f"發證單位：{org}"], font_size=18)
            if cert_file:
                add_image(s, cert_file, 7.5, 1.6, 4)

    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_ending(s)
    add_footer(prs, FOOTER_TEXT)
    path = os.path.join(OUTPUT_DIR, "C-7a_專業證照.pptx")
    prs.save(path)
    print(f"[OK] C-7a: {os.path.getsize(path):,} bytes")


def build_cert_other_ppt():
    """C-7b: 其他證照（語言/資訊等）"""
    prs = new_presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_cover(s, "C-7 檢定證照", "其他檢定證照")

    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_toc(s, [("01", "其他證照", "")])

    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_ending(s)
    add_footer(prs, FOOTER_TEXT)
    path = os.path.join(OUTPUT_DIR, "C-7b_其他證照.pptx")
    prs.save(path)
    print(f"[OK] C-7b: {os.path.getsize(path):,} bytes")


def build_achievement_ppt():
    """C-8: 特殊優良表現證明"""
    prs = new_presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_cover(s, "C-8 特殊優良表現證明")

    toc_items = []
    for i, a in enumerate(ACTIVITIES):
        if a["name"] and a["name"].startswith("{{") == False:
            toc_items.append((f"{i+1:02d}", a["name"], a.get("org", "")))
    add_toc(s, toc_items if toc_items else [("01", "課外活動與研習", "")])

    for a in ACTIVITIES:
        if a["name"] and a["name"].startswith("{{") == False:
            s = prs.slides.add_slide(prs.slide_layouts[6])
            add_title(s, a["name"])
            lines = []
            if a.get("org") and a["org"].startswith("{{") == False:
                lines.append(f"主辦單位：{a['org']}")
            if a.get("hours") and a["hours"].startswith("{{") == False:
                lines.append(f"時數：{a['hours']}")
            if a.get("learn") and a["learn"].startswith("{{") == False:
                lines.append("")
                lines.append(a["learn"])
            add_multiline_text(s, 0.8, 1.6, 6, 4, lines, font_size=18)
            if a.get("cert") and a["cert"]:
                add_image(s, a["cert"], 7.5, 1.6, 4)

    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_ending(s)
    add_footer(prs, FOOTER_TEXT)
    path = os.path.join(OUTPUT_DIR, "C-8_特殊優良表現.pptx")
    prs.save(path)
    print(f"[OK] C-8: {os.path.getsize(path):,} bytes")


def build_overview_ppt():
    """D-1: 多元表現綜整心得"""
    prs = new_presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_cover(s, "D-1 多元表現綜整心得")

    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_toc(s, [
        ("01", "多元表現清冊", ""),
        ("02", "能力分析", ""),
        ("03", "總結", ""),
    ])

    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(s, "多元表現清冊")
    add_multiline_text(s, 0.8, 1.6, 5.5, 5, [
        f"姓名：{STUDENT_NAME_ZH}",
        f"學校：{STUDENT_SCHOOL}{STUDENT_DEPT}",
        "",
        "競賽：",
        f"• {AWARD_1_NAME}：{AWARD_1_RANK}",
        "",
        "檢定證照：",
        f"• {CERT_1_NAME}",
        f"• {CERT_2_NAME}",
    ], font_size=16)

    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(s, "總結")
    if SUMMARY_TEXT and SUMMARY_TEXT.startswith("{{") == False:
        add_multiline_text(s, 0.8, 2, 11.5, 4, [SUMMARY_TEXT], font_size=18)
    else:
        add_multiline_text(s, 0.8, 2, 11.5, 4, [
            "三年的學習歷程，建立從設計到製造的完整思維。",
            "跨領域整合機械、電子、軟體能力。",
        ], font_size=18)

    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_ending(s)
    add_footer(prs, FOOTER_TEXT)
    path = os.path.join(OUTPUT_DIR, "D-1_綜整心得.pptx")
    prs.save(path)
    print(f"[OK] D-1: {os.path.getsize(path):,} bytes")


def build_motivation_ppt():
    """D-2: 學習歷程自述（就讀動機 + 未來規劃）"""
    prs = new_presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_cover(s, "D-2 學習歷程自述", f"{STUDENT_NAME_ZH} · {STUDENT_SCHOOL}{STUDENT_DEPT}")

    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_toc(s, [
        ("01", "學習歷程", "從入學到專題的技術累積"),
        ("02", "就讀動機", f"為什麼選擇{TARGET_SCHOOL}{TARGET_DEPT}"),
        ("03", "未來規劃", "短中長期目標"),
    ])

    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(s, "學習歷程")
    add_multiline_text(s, 0.8, 1.6, 11.5, 5.5, [
        f"高一：基礎訓練 — 手繪製圖、車床、銑床、鉗工",
        f"高二：專業深化 — AutoCAD、Inventor、基本電學",
        f"高三：跨域整合 — {PROJECT_NAME}",
    ], font_size=18)

    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(s, "就讀動機")
    lines = [MOTIVATION_STATEMENT]
    if KEY_LABELS:
        lines.append("")
        lines.append("三個關鍵能力：")
        for lbl in KEY_LABELS:
            if lbl and lbl.startswith("{{") == False:
                lines.append(f"• {lbl}")
    add_multiline_text(s, 0.8, 1.6, 5.5, 5, lines, font_size=18)

    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(s, "未來規劃")
    plan_lines = ["短期目標：", COLLEGE_PLAN, "", "長期目標：", LONG_TERM_GOAL]
    add_multiline_text(s, 0.8, 1.6, 5.5, 5, plan_lines, font_size=18)

    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_ending(s)
    add_footer(prs, FOOTER_TEXT)
    path = os.path.join(OUTPUT_DIR, "D-2_學習歷程自述.pptx")
    prs.save(path)
    print(f"[OK] D-2: {os.path.getsize(path):,} bytes")


def build_other_ppt():
    """D-3: 其他有利審查資料（GitHub 專案、職業訪談等）"""
    prs = new_presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_cover(s, "D-3 其他有利審查資料")

    s = prs.slides.add_slide(prs.slide_layouts[6])
    toc_items = [("01", proj["name"], proj["tech"]) for proj in OTHER_PROJECTS if proj["name"] and proj["name"].startswith("{{") == False]
    add_toc(s, toc_items if toc_items else [("01", "其他作品", "")])

    for proj in OTHER_PROJECTS:
        if proj["name"] and proj["name"].startswith("{{") == False:
            s = prs.slides.add_slide(prs.slide_layouts[6])
            add_title(s, proj["name"])
            lines = []
            if proj.get("tech") and proj["tech"].startswith("{{") == False:
                lines.append(f"技術：{proj['tech']}")
            if proj.get("link") and proj["link"].startswith("{{") == False:
                lines.append(f"連結：{proj['link']}")
            if proj.get("contrib") and proj["contrib"].startswith("{{") == False:
                lines.append("")
                lines.append(f"貢獻：{proj['contrib']}")
            add_multiline_text(s, 0.8, 1.6, 11.5, 4, lines, font_size=18)

    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_ending(s)
    add_footer(prs, FOOTER_TEXT)
    path = os.path.join(OUTPUT_DIR, "D-3_其他有利審查.pptx")
    prs.save(path)
    print(f"[OK] D-3: {os.path.getsize(path):,} bytes")


# ====================================================================
# 主程式 — Agent 請根據使用者需要的類別啟用對應函式
# ====================================================================
if __name__ == "__main__":
    print("=== 開始產生備審 PPT ===\n")

    # B 類
    build_project_ppt()        # B-1a 專題實作
    # build_internship_ppt()   # B-1b 實習科目（如有需要請取消註解）
    # build_handdraft_ppt()    # B-2a 手繪製圖
    # build_cad_ppt()          # B-2b 電腦輔助製圖
    # build_3dmodel_ppt()      # B-2c 3D 建模

    # C 類
    # build_flexible_ppt()     # C-1 彈性學習
    build_competition_ppt()    # C-5 競賽表現
    build_cert_pro_ppt()       # C-7a 專業證照
    # build_cert_other_ppt()   # C-7b 其他證照
    # build_achievement_ppt()  # C-8 特殊優良表現

    # D 類
    # build_overview_ppt()     # D-1 綜整心得
    build_motivation_ppt()     # D-2 學習歷程自述
    # build_other_ppt()        # D-3 其他有利審查

    print(f"\n=== 完成！檔案位於: {OUTPUT_DIR} ===")
